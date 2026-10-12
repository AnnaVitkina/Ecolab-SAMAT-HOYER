"""Detect and normalize Samat Tesjoki route-rate sheets."""

from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from ebs_format import parse_rate_value

SIGNATURE_MARKERS = {"route", "unloading city"}

VENDOR_NAME = "Samat Haanpa - Tesjoki"
ORIGIN_COUNTRY = "FI"
ORIGIN_POSTAL_CODE = "07"
ORIGIN_CITY = "Tesjoki"
DEFAULT_MODE = "Road"

ROUTE_PATTERN = re.compile(
    r"^TEJ-(?P<country>[A-Z]{2})-(?P<postal>.+)$",
    re.IGNORECASE,
)


def find_tesjoki_header_row(raw: pd.DataFrame, *, max_scan: int = 40) -> int | None:
    limit = min(len(raw), max_scan)
    for idx in range(limit):
        values = {
            str(value).strip().casefold()
            for value in raw.iloc[idx].tolist()
            if pd.notna(value) and str(value).strip()
        }
        if SIGNATURE_MARKERS.issubset(values):
            return int(idx)
    return None


def _sheet_has_tesjoki_headers(path: Path, sheet_name: str) -> bool:
    preview = pd.read_excel(
        path,
        sheet_name=sheet_name,
        header=None,
        nrows=40,
        engine="openpyxl",
    )
    return find_tesjoki_header_row(preview) is not None


def find_tesjoki_sheet(path: Path) -> str | None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet_name in workbook.sheetnames:
            if _sheet_has_tesjoki_headers(path, sheet_name):
                return sheet_name
    finally:
        workbook.close()
    return None


def is_tesjoki_path(path: Path) -> bool:
    if path.suffix.lower() != ".xlsx":
        return False
    return find_tesjoki_sheet(path) is not None


def is_tesjoki_format(df: pd.DataFrame) -> bool:
    columns = {str(column).strip().casefold() for column in df.columns}
    return SIGNATURE_MARKERS.issubset(columns)


def load_tesjoki(path: Path) -> pd.DataFrame:
    sheet_name = find_tesjoki_sheet(path)
    if sheet_name is None:
        raise ValueError(f"Tesjoki rate headers not found in {path.name}")

    raw = pd.read_excel(path, sheet_name=sheet_name, header=None, engine="openpyxl")
    header_row = find_tesjoki_header_row(raw)
    if header_row is None:
        raise ValueError(f"Tesjoki rate headers not found in sheet '{sheet_name}'")

    columns = [
        str(value).strip() if pd.notna(value) and str(value).strip() else f"Unnamed: {idx}"
        for idx, value in enumerate(raw.iloc[header_row].tolist())
    ]
    data = raw.iloc[header_row + 1 :].copy()
    data.columns = columns
    data = data.dropna(how="all").reset_index(drop=True)
    return data


def _cell_string(value: object) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    text = str(value).strip()
    if text.casefold() in {"", "nan", "none", "nat"}:
        return ""
    return text


def _find_column(columns: pd.Index, *needles: str) -> str | None:
    for column in columns:
        text = str(column).strip().casefold()
        if any(needle in text for needle in needles):
            return column
    return None


def _parse_route(route: object) -> tuple[str, str] | None:
    text = _cell_string(route)
    match = ROUTE_PATTERN.match(text)
    if not match:
        return None
    country = match.group("country").upper()
    postal = match.group("postal").strip()
    if postal.endswith(".0"):
        postal = postal[:-2]
    return country, postal


def _select_rate_column(columns: pd.Index) -> str:
    rate_columns = [
        column
        for column in columns
        if "rate" in str(column).strip().casefold()
    ]
    if not rate_columns:
        raise ValueError(
            "Tesjoki sheet is missing a rate column "
            "(expected a column containing 'Rates')."
        )

    preferred = [
        column
        for column in rate_columns
        if "2026" in str(column) or "june" in str(column).casefold()
    ]
    return preferred[-1] if preferred else rate_columns[-1]


def normalize_tesjoki_df(df: pd.DataFrame) -> pd.DataFrame:
    route_col = _find_column(df.columns, "route")
    city_col = _find_column(df.columns, "unloading city")
    if route_col is None:
        raise ValueError("Tesjoki sheet is missing required column: ROUTE")

    rate_col = _select_rate_column(df.columns)

    rows: list[dict[str, object]] = []
    for _, source_row in df.iterrows():
        parsed = _parse_route(source_row[route_col])
        if parsed is None:
            continue
        destination_country, destination_postal = parsed
        freight_rate = parse_rate_value(source_row[rate_col])
        if freight_rate is None:
            continue

        row: dict[str, object] = {
            "Lane": _cell_string(source_row[route_col]),
            "Origin Country": ORIGIN_COUNTRY,
            "Origin ZIP Code": ORIGIN_POSTAL_CODE,
            "Origin City": ORIGIN_CITY,
            "Destination Country": destination_country,
            "Destination ZIP Code": destination_postal,
            "Vendor Name": VENDOR_NAME,
            "Mode": DEFAULT_MODE,
            "Freight Rate": freight_rate,
        }
        if city_col:
            row["Destination City"] = _cell_string(source_row[city_col])
        rows.append(row)

    if not rows:
        raise ValueError("No valid Tesjoki route/rate rows found.")

    return pd.DataFrame(rows)
